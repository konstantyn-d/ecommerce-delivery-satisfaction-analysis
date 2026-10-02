# %% [markdown]
# # 06 - Results of the review coding (Phase 5)
#
# Goal: turn the manually coded reviews into a results table.
#
# Input : review_coding/coding_sheet.xlsx   (coded by hand)
#         review_coding/recode_sheet.xlsx   (optional: the 30 reviews coded a second time)
#         review_coding/sample_key.csv      (delivery information, joined here for the first time)
#         review_coding/categories.csv      (the valid categories)
# Output: review_coding/coded_reviews.csv   (the coding as a plain, version-controlled file)
#         review_coding/results.md          (generated - do not edit by hand)
#
# The script stops with a clear message if the coding is incomplete or contains an unknown category.
#
# Usage: python notebooks/06_review_coding_results.py [folder]
#        folder = where the coding files are; default: review_coding/

# %%
import csv
import math
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

# Windows consoles default to a legacy code page; force UTF-8 so Portuguese text prints safely.
sys.stdout.reconfigure(encoding="utf-8")

try:
    ROOT = Path(__file__).resolve().parents[1]
except NameError:  # running cell by cell: fall back to the working directory
    ROOT = Path.cwd() if (Path.cwd() / "sql").exists() else Path.cwd().parent

FOLDER = Path(sys.argv[1]) if len(sys.argv) > 1 and not sys.argv[1].startswith("-") else ROOT / "review_coding"


# %% [markdown]
# ## 1. Read and validate the coding

# %%
def read_sheet(path: Path) -> list:
    """Read the 'coding' tab of an Excel coding sheet as a list of dictionaries."""
    sheet = load_workbook(path, read_only=True, data_only=True)["coding"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, [clean(v) for v in row])) for row in rows[1:] if row[0] is not None]


def clean(value):
    """Trim text cells; turn empty cells into None."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_csv(path: Path) -> list:
    with open(path, encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


CATEGORIES = [row["category"] for row in read_csv(FOLDER / "categories.csv")]
coded = read_sheet(FOLDER / "coding_sheet.xlsx")
key = {int(row["sample_id"]): row for row in read_csv(FOLDER / "sample_key.csv")}

problems = []
for row in coded:
    sid = row["sample_id"]
    if row["primary_category"] is None:
        problems.append(f"sample_id {sid}: primary_category is empty")
    elif row["primary_category"] not in CATEGORIES:
        problems.append(f"sample_id {sid}: unknown primary_category '{row['primary_category']}'")
    if row["secondary_category"] is not None and row["secondary_category"] not in CATEGORIES:
        problems.append(f"sample_id {sid}: unknown secondary_category '{row['secondary_category']}'")
    if row["secondary_category"] is not None and row["secondary_category"] == row["primary_category"]:
        problems.append(f"sample_id {sid}: secondary_category repeats the primary_category")
if len(coded) != len(key):
    problems.append(f"the sheet has {len(coded)} rows, the sample has {len(key)}")

if problems:
    print(f"The coding is not ready: {len(problems)} problem(s).")
    for line in problems[:25]:
        print("  -", line)
    if len(problems) > 25:
        print(f"  ... and {len(problems) - 25} more")
    sys.exit(1)

# join the delivery information (kept apart until now, so the coding was blind)
for row in coded:
    info = key[int(row["sample_id"])]
    row["review_id"] = info["review_id"]
    row["delivery_group"] = info["delivery_group"]
    row["delivery_outcome"] = info["delivery_outcome"]

with open(FOLDER / "coded_reviews.csv", "w", newline="", encoding="utf-8-sig") as handle:
    columns = ["sample_id", "review_id", "review_score", "primary_category", "secondary_category",
               "good_quote", "notes", "delivery_outcome", "delivery_group"]
    writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(coded)


# %% [markdown]
# ## 2. Helper functions

# %%
def wilson(successes: int, total: int) -> tuple:
    """95% Wilson confidence interval for a proportion."""
    z = 1.96
    p = successes / total
    centre = (p + z * z / (2 * total)) / (1 + z * z / total)
    half = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total)) / (1 + z * z / total)
    return (max(0.0, centre - half), min(1.0, centre + half))


def md_table(headers, rows) -> list:
    lines = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    lines += ["| " + " | ".join(str(v) for v in row) + " |" for row in rows]
    return lines + [""]


def cohen_kappa(first: list, second: list) -> float:
    """Agreement between two codings of the same items, corrected for chance agreement."""
    total = len(first)
    observed = sum(a == b for a, b in zip(first, second)) / total
    count_1, count_2 = Counter(first), Counter(second)
    expected = sum(count_1[c] * count_2[c] for c in set(first) | set(second)) / (total * total)
    return 1.0 if expected == 1 else (observed - expected) / (1 - expected)


def share_columns(rows_by_group: dict, groups: list) -> list:
    """Category shares inside each group (column percentages), with the number of reviews."""
    body = []
    for category in CATEGORIES:
        line = [category]
        for group in groups:
            members = rows_by_group[group]
            hits = sum(r["primary_category"] == category for r in members)
            line.append(f"{hits / len(members):.0%} ({hits})")
        body.append(line)
    body.append(["**Reviews**"] + [f"**{len(rows_by_group[g])}**" for g in groups])
    return body


# %% [markdown]
# ## 3. Results

# %%
total = len(coded)
primary = Counter(r["primary_category"] for r in coded)
mentioned = Counter()
for r in coded:
    mentioned.update({r["primary_category"], r["secondary_category"]} - {None})

report = [
    "# Review coding results",
    "",
    "Generated by `notebooks/06_review_coding_results.py` from `coding_sheet.xlsx`. "
    "Do not edit by hand: re-run the script instead.",
    "",
    f"Sample: {total} reviews with 1-2 stars and a written comment "
    "([`sampling_report.md`](sampling_report.md)). Categories: [`codebook.md`](codebook.md). "
    "The reviews were coded without knowing whether the order was late.",
    "",
    "## 1. What customers complain about",
    "",
    "*Primary* = the main complaint of the review (each review counted once). "
    "*Mentioned* = primary or secondary complaint. The confidence interval (Wilson, 95%) belongs to the primary share.",
    "",
]
report += md_table(
    ["category", "reviews (primary)", "share", "95% CI", "reviews (mentioned)", "share mentioned"],
    [
        (c, primary[c], f"{primary[c] / total:.1%}",
         "{:.0%} - {:.0%}".format(*wilson(primary[c], total)),
         mentioned[c], f"{mentioned[c] / total:.1%}")
        for c in sorted(CATEGORIES, key=lambda c: primary[c], reverse=True)
    ],
)

outcomes = ["On time", "Late", "Not delivered"]
by_outcome = {o: [r for r in coded if r["delivery_outcome"] == o] for o in outcomes}
report += [
    "## 2. By delivery outcome",
    "",
    "Share of each primary category inside the group (number of reviews in brackets). "
    "Groups are small: read differences of less than about 15 points as noise.",
    "",
]
report += md_table(["category"] + outcomes, share_columns(by_outcome, outcomes))

groups = ["On time, single item", "On time, several items", "Late", "Not delivered"]
by_group = {g: [r for r in coded if r["delivery_group"] == g] for g in groups}
report += ["## 3. On-time orders split by order size", ""]
report += md_table(["category"] + groups, share_columns(by_group, groups))

# %% [markdown]
# ## 4. Consistency check: the 30 reviews coded a second time

# %%
report += ["## 4. Consistency check (intra-coder reliability)", ""]
recode_path = FOLDER / "recode_sheet.xlsx"
recoded = [r for r in read_sheet(recode_path) if r["primary_category"] is not None] if recode_path.exists() else []
if not recoded:
    report += ["Not done yet: `recode_sheet.xlsx` has no coded rows. [TBD]", ""]
else:
    first_by_id = {int(r["sample_id"]): r["primary_category"] for r in coded}
    unknown = [r for r in recoded if r["primary_category"] not in CATEGORIES]
    if unknown:
        sys.exit(f"recode_sheet.xlsx contains {len(unknown)} unknown categories - fix them and re-run.")
    first = [first_by_id[int(r["sample_id"])] for r in recoded]
    second = [r["primary_category"] for r in recoded]
    agreement = sum(a == b for a, b in zip(first, second)) / len(recoded)
    report += md_table(
        ["measure", "value"],
        [
            ("Reviews coded twice", len(recoded)),
            ("Same primary category both times", f"{sum(a == b for a, b in zip(first, second))} ({agreement:.0%})"),
            ("Cohen's kappa", f"{cohen_kappa(first, second):.2f}"),
        ],
    )
    disagreements = [(r["sample_id"], a, b) for r, a, b in zip(recoded, first, second) if a != b]
    if disagreements:
        report += ["Disagreements:", ""]
        report += md_table(["sample_id", "first coding", "second coding"], disagreements)

# %% [markdown]
# ## 5. Quote candidates (Portuguese originals, to be translated for the deck)

# %%
quotes = [r for r in coded if r["good_quote"] is not None]
report += ["## 5. Quote candidates", ""]
if not quotes:
    report += ["No review was marked in the `good_quote` column. [TBD]", ""]
else:
    report += md_table(
        ["sample_id", "category", "delivery outcome", "original comment (Portuguese)"],
        [(r["sample_id"], r["primary_category"], r["delivery_outcome"],
          " ".join(str(r["comment_message"]).split()).replace("|", "/")) for r in quotes],
    )

(FOLDER / "results.md").write_text("\n".join(report), encoding="utf-8")
print("\n".join(report))
print(f"\nwritten: {FOLDER / 'results.md'}\nwritten: {FOLDER / 'coded_reviews.csv'}")

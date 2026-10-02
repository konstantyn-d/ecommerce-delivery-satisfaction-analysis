# %% [markdown]
# # 05 - Review sample for qualitative coding (Phase 5)
#
# Goal: draw a reproducible random sample of 300 bad reviews (1-2 stars, with a written comment)
# and prepare the files for manual coding.
#
# Input : data/processed/olist.duckdb, review_coding/categories.csv
# Output: review_coding/sample.csv            the 300 reviews to code (no delivery information: blind coding)
#         review_coding/sample_key.csv        delivery information per review, joined AFTER coding
#         review_coding/pilot.csv             60 other reviews, used only to develop the codebook
#         review_coding/coding_sheet.xlsx     Excel sheet with drop-down lists for the coding
#         review_coding/recode_sheet.xlsx     30 of the 300 reviews, to be coded a second time later
#         review_coding/sampling_report.md    how the sample was drawn (generated)
#
# The two Excel sheets are only created if they do not exist yet, so that re-running this script
# can never overwrite coding work that has already been done.

# %%
import csv
import math
import sys
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# Windows consoles default to a legacy code page; force UTF-8 so Portuguese text prints safely.
sys.stdout.reconfigure(encoding="utf-8")

try:
    ROOT = Path(__file__).resolve().parents[1]
except NameError:  # running cell by cell: fall back to the working directory
    ROOT = Path.cwd() if (Path.cwd() / "sql").exists() else Path.cwd().parent

DB_PATH = ROOT / "data" / "processed" / "olist.duckdb"
OUT = ROOT / "review_coding"

SAMPLE_SIZE = 300
PILOT_PER_GROUP = 15
RECODE_SIZE = 30
SEED = "20261002"  # fixed text added to every id before hashing: same seed -> same sample, forever

con = duckdb.connect(str(DB_PATH), read_only=True)

# %% [markdown]
# ## 1. Sampling frame
# Every review that could be selected: score 1-2, a written comment, order in the analysis period.
# One row per review_id: a review that covers several orders must not be coded twice (decision #12).
# The delivery group is the stratum: the sample gets the same mix of groups as the frame.

# %%
con.execute(
    """
    CREATE TEMP VIEW frame AS
    SELECT r.review_id,
           r.order_id,
           r.review_score,
           raw.review_comment_title,
           r.review_comment_message,
           CASE WHEN f.delivery_outcome = 'On time' AND f.item_count >= 2 THEN 'On time, several items'
                WHEN f.delivery_outcome = 'On time'                      THEN 'On time, single item'
                ELSE f.delivery_outcome END AS delivery_group,
           f.delivery_outcome, f.delay_band, f.delay_days, f.item_count, f.seller_count,
           f.review_before_delivery, f.customer_state
    FROM stg.reviews r
    JOIN model.fact_orders f USING (order_id)
    JOIN raw.order_reviews raw ON raw.review_id = r.review_id AND raw.order_id = r.order_id
    WHERE r.review_score <= 2
      AND r.review_comment_message IS NOT NULL
      AND f.delivery_outcome <> 'Unknown'            -- 1 review; too few to form a group
    QUALIFY row_number() OVER (PARTITION BY r.review_id ORDER BY r.order_id) = 1
    """
)
frame_size = con.execute("SELECT count(*) FROM frame").fetchone()[0]
groups = con.execute("SELECT delivery_group, count(*) FROM frame GROUP BY 1 ORDER BY 2 DESC").fetchall()

# %% [markdown]
# ## 2. How many reviews per group? Proportional allocation
# Each group gets the same share of the sample as it has in the frame. The exact shares are not whole
# numbers, so they are rounded down and the remaining places go to the groups with the largest remainders.

# %%
exact = {g: SAMPLE_SIZE * size / frame_size for g, size in groups}
allocation = {g: math.floor(x) for g, x in exact.items()}
left_over = SAMPLE_SIZE - sum(allocation.values())
for g in sorted(exact, key=lambda g: exact[g] - allocation[g], reverse=True)[:left_over]:
    allocation[g] += 1

# %% [markdown]
# ## 3. Draw the sample
# Random but reproducible: every review gets a pseudo-random position from md5(review_id + seed).
# Inside each group the first n positions are the sample; the next 15 are the pilot set.
# No random number generator is involved, so the result is identical on every machine.

# %%
con.execute(
    f"""
    CREATE TEMP TABLE ranked AS
    SELECT *, row_number() OVER (PARTITION BY delivery_group ORDER BY md5(review_id || '{SEED}')) AS position
    FROM frame
    """
)
con.execute("CREATE TEMP TABLE allocation (delivery_group VARCHAR, n INTEGER)")
con.executemany("INSERT INTO allocation VALUES (?, ?)", list(allocation.items()))

# sample_id 1..300 in a shuffled order, so that the groups are mixed in the coding sheet
con.execute(
    f"""
    CREATE TEMP TABLE sample AS
    SELECT row_number() OVER (ORDER BY md5(r.review_id || '{SEED}' || 'order')) AS sample_id, r.*
    FROM ranked r JOIN allocation a USING (delivery_group)
    WHERE r.position <= a.n
    """
)
con.execute(
    f"""
    CREATE TEMP TABLE pilot AS
    SELECT r.*
    FROM ranked r JOIN allocation a USING (delivery_group)
    WHERE r.position > a.n AND r.position <= a.n + {PILOT_PER_GROUP}
    """
)
# the 30 reviews to code a second time: again chosen by hash, in a different order
con.execute(
    f"""
    CREATE TEMP TABLE recode AS
    SELECT row_number() OVER (ORDER BY md5(review_id || '{SEED}' || 'recode order')) AS recode_id, *
    FROM (SELECT * FROM sample ORDER BY md5(review_id || '{SEED}' || 'recode') LIMIT {RECODE_SIZE})
    """
)


# %% [markdown]
# ## 4. Write the files

# %%
def write_csv(path: Path, sql: str) -> None:
    """Write a query result as CSV. utf-8-sig lets Excel show Portuguese accents correctly."""
    cursor = con.execute(sql)
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([col[0] for col in cursor.description])
        writer.writerows(cursor.fetchall())
    print(f"written: {path.relative_to(ROOT)}")


write_csv(
    OUT / "sample.csv",
    "SELECT sample_id, review_id, review_score, review_comment_title, review_comment_message FROM sample ORDER BY sample_id",
)
write_csv(
    OUT / "sample_key.csv",
    """SELECT sample_id, review_id, order_id, delivery_group, delivery_outcome, delay_band, delay_days,
              item_count, seller_count, review_before_delivery, customer_state
       FROM sample ORDER BY sample_id""",
)
write_csv(
    OUT / "pilot.csv",
    """SELECT delivery_group, review_id, review_score, review_comment_title, review_comment_message
       FROM pilot ORDER BY delivery_group, position""",
)

# %%
with open(OUT / "categories.csv", encoding="utf-8") as handle:
    categories = list(csv.DictReader(handle))

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
WRAP = Alignment(wrap_text=True, vertical="top")


def build_sheet(path: Path, id_column: str, sql: str, with_extras: bool) -> None:
    """Create an Excel coding sheet: one review per row, drop-down lists for the categories."""
    if path.exists():
        print(f"kept (already exists, not overwritten): {path.relative_to(ROOT)}")
        return

    book = Workbook()
    sheet = book.active
    sheet.title = "coding"
    headers = [id_column, "review_score", "comment_title", "comment_message", "primary_category"]
    if with_extras:
        headers += ["secondary_category", "good_quote", "notes"]
    sheet.append(headers)
    for row in con.execute(sql).fetchall():
        sheet.append(list(row))

    # second tab: the category list (source of the drop-downs and a quick reference while coding)
    reference = book.create_sheet("categories")
    reference.append(["category", "short_definition"])
    for item in categories:
        reference.append([item["category"], item["short_definition"]])
    reference.column_dimensions["A"].width = 34
    reference.column_dimensions["B"].width = 120

    last_row = sheet.max_row
    dropdown = DataValidation(type="list", formula1=f"=categories!$A$2:$A${len(categories) + 1}", allow_blank=True)
    dropdown.error = "Choose a category from the list."
    dropdown.showErrorMessage = True
    sheet.add_data_validation(dropdown)
    dropdown.add(f"E2:E{last_row}")
    if with_extras:
        dropdown.add(f"F2:F{last_row}")

    widths = {"A": 10, "B": 8, "C": 24, "D": 95, "E": 30, "F": 30, "G": 11, "H": 40}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
        text_length = len(str(row[3].value or ""))
        sheet.row_dimensions[row[0].row].height = 16 * max(1, math.ceil(text_length / 95))
    sheet.freeze_panes = "A2"

    book.save(path)
    print(f"written: {path.relative_to(ROOT)}")


build_sheet(
    OUT / "coding_sheet.xlsx", "sample_id",
    "SELECT sample_id, review_score, review_comment_title, review_comment_message FROM sample ORDER BY sample_id",
    with_extras=True,
)
build_sheet(
    OUT / "recode_sheet.xlsx", "sample_id",
    "SELECT sample_id, review_score, review_comment_title, review_comment_message FROM recode ORDER BY recode_id",
    with_extras=False,
)

# %%
check = con.execute("SELECT count(*), count(DISTINCT review_id), count(DISTINCT order_id) FROM sample").fetchone()
overlap = con.execute("SELECT count(*) FROM sample s JOIN pilot p USING (review_id)").fetchone()[0]
score_split = con.execute("SELECT review_score, count(*) FROM sample GROUP BY 1 ORDER BY 1").fetchall()
frame_scores = dict(con.execute("SELECT review_score, count(*) FROM frame GROUP BY 1").fetchall())

report = [
    "# Sampling report",
    "",
    "Generated by `notebooks/05_review_sample.py`. Do not edit by hand: re-run the script instead.",
    "",
    "## Method",
    "",
    "- **Frame:** reviews with a score of 1 or 2 and a written comment, on orders purchased from 2017-01-01 to "
    f"2018-08-31, one row per `review_id`: **{frame_size:,} reviews**.",
    f"- **Sample size:** {SAMPLE_SIZE} reviews.",
    "- **Design:** stratified random sample with proportional allocation. Strata = delivery group of the order. "
    "Because each group has the same share in the sample as in the frame, the overall results need no weighting.",
    f"- **Randomisation:** reviews are sorted by `md5(review_id || '{SEED}')` inside each group and the first "
    "*n* are taken. The sample is therefore identical on every run and on every machine.",
    "- **Blind coding:** the coding sheet contains only the score and the text. The delivery group is kept in "
    "`sample_key.csv` and joined after coding.",
    f"- **Pilot set:** the next {PILOT_PER_GROUP} reviews of each group ({PILOT_PER_GROUP * len(groups)} in total), "
    "used to develop the codebook. None of them is in the sample.",
    f"- **Re-coding set:** {RECODE_SIZE} of the {SAMPLE_SIZE} sampled reviews, selected by hash, to measure "
    "whether the coder is consistent with himself.",
    "",
    "## Allocation",
    "",
    "| delivery group | reviews in frame | share of frame | reviews in sample | share of sample |",
    "|---|---|---|---|---|",
]
sample_counts = dict(con.execute("SELECT delivery_group, count(*) FROM sample GROUP BY 1").fetchall())
for g, size in groups:
    report.append(
        f"| {g} | {size:,} | {size / frame_size:.1%} | {sample_counts[g]} | {sample_counts[g] / SAMPLE_SIZE:.1%} |"
    )
report += [
    f"| **Total** | **{frame_size:,}** | 100.0% | **{sum(sample_counts.values())}** | 100.0% |",
    "",
    "## Checks",
    "",
    "| check | result |",
    "|---|---|",
    f"| rows in the sample | {check[0]} |",
    f"| distinct review_id in the sample | {check[1]} |",
    f"| distinct order_id in the sample | {check[2]} |",
    f"| reviews that are both in the sample and in the pilot set | {overlap} |",
]
for score, count in score_split:
    report.append(
        f"| sample reviews with {score} star(s) | {count} ({count / SAMPLE_SIZE:.1%}; "
        f"frame: {frame_scores[score] / frame_size:.1%}) |"
    )
report.append("")

(OUT / "sampling_report.md").write_text("\n".join(report), encoding="utf-8")
print("\n".join(report))
con.close()
